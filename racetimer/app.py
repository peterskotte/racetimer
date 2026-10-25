"""RaceTimer – FastAPI-server der binder læsere, database, motor og skærme sammen.

Kør:  uvicorn racetimer.app:app --host 0.0.0.0 --port 8000

Normalt kører alt på Pi'en. Med RACETIMER_MODE=mirror kører den i stedet som
offentlig spejlserver på en webserver: ingen læsere, ingen admin – den viser kun
det, Pi'en sender til /api/mirror/push (se sync.py).
"""
import asyncio
import csv
import gzip
import hashlib
import secrets
import io
import json
import os
import re
import threading
import time

from fastapi import Depends, FastAPI, File, HTTPException, Request, Response, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles

from . import branding, db, engine, events, modules, outcome, readers, signup, sync, updater

HERE = os.path.dirname(__file__)
STATIC = os.path.join(HERE, "static")
READ_LOG = os.path.join(os.path.dirname(os.path.abspath(db.DB_PATH)), "reads.log")  # rå backup ved siden af databasen
MIRROR = os.environ.get("RACETIMER_MODE") == "mirror"


def _code_version():
    """Fingeraftryk af programmets filer – skifter ved hver opdatering, så skærmene genindlæser sig selv."""
    h = hashlib.sha1()
    for dirpath, _, files in sorted(os.walk(HERE)):
        for f in sorted(files):
            if f.endswith((".py", ".html", ".css", ".js", ".png", ".webp")):
                with open(os.path.join(dirpath, f), "rb") as fh:
                    h.update(f.encode() + fh.read())
    return h.hexdigest()[:12]


APP_VERSION = _code_version()

_host_cache = {"t": 0, "v": None}


def host_info():
    """Pi'ens navn og aktuelle IP – så adressen kan aflæses på skærmen, også efter et WiFi-skift."""
    import socket
    import subprocess
    if time.time() - _host_cache["t"] < 10:
        return _host_cache["v"]
    ip = None
    try:  # IP'en på den forbindelse, der bruges ud mod nettet (der sendes ingen pakker)
        with socket.socket(socket.AF_INET, socket.SOCK_DGRAM) as sk:
            sk.connect(("10.255.255.255", 1))
            ip = sk.getsockname()[0]
    except OSError:
        pass
    if not ip or ip.startswith("127."):
        try:  # fx Pi'ens eget hotspot uden internet
            ip = (subprocess.run(["hostname", "-I"], capture_output=True, text=True, timeout=2).stdout.split() or [None])[0]
        except (OSError, subprocess.SubprocessError):
            ip = None
    name = socket.gethostname().split(".")[0]
    _host_cache.update(t=time.time(), v={"ip": ip if ip and not ip.startswith("127.") else None, "name": f"{name}.local",
                                         "wifi": wifi_ssid()})
    return _host_cache["v"]


def wifi_ssid():
    """Navnet på det WiFi, Pi'en er forbundet til (None ved kabel eller uden NetworkManager)."""
    import re
    import subprocess
    env = dict(os.environ, LC_ALL="C")  # på en dansk Pi skriver nmcli ellers "ja" i stedet for "yes"
    try:
        out = subprocess.run([os.environ.get("RACETIMER_NMCLI", "nmcli"), "-t", "-f", "ACTIVE,SSID", "device", "wifi",
                              "list", "--rescan", "no"], capture_output=True, text=True, timeout=3, env=env).stdout
    except (OSError, subprocess.SubprocessError):
        out = ""
    for line in out.splitlines():
        active, _, ssid = line.partition(":")  # nmcli -t escaper ':' i navnet som '\:'
        if active in ("yes", "ja") and ssid:
            return re.sub(r"\\(.)", r"\1", ssid)
    try:  # reserve uden NetworkManager
        return subprocess.run(["iwgetid", "-r"], capture_output=True, text=True, timeout=2, env=env).stdout.strip() or None
    except (OSError, subprocess.SubprocessError):
        return None


_clock_cache = {"t": 0, "v": None}


def clock_info():
    """Er Pi'ens ur stillet via internettet (NTP) siden opstart? None = kan ikke afgøres (fx på en Mac).

    Uden RTC-batteri starter Pi'en med det klokkeslæt, den havde ved sidste slukning, indtil den får net.
    """
    import subprocess
    if time.time() - _clock_cache["t"] < 10:
        return _clock_cache["v"]
    forced = os.environ.get("RACETIMER_CLOCK_SYNCED")  # kun til test
    synced = None
    if forced in ("0", "1"):
        synced = forced == "1"
    elif os.path.exists("/run/systemd/timesync/synchronized"):
        synced = True  # systemd-timesyncd har stillet uret siden opstart
    else:
        try:
            out = subprocess.run(["timedatectl", "show", "-p", "NTPSynchronized", "--value"],
                                 capture_output=True, text=True, timeout=2).stdout.strip()
            synced = {"yes": True, "no": False}.get(out)
        except (OSError, subprocess.SubprocessError):
            pass
    _clock_cache.update(t=time.time(), v={"synced": synced})
    return _clock_cache["v"]

app = FastAPI(title="RaceTimer")
app.mount("/static", StaticFiles(directory=STATIC), name="static")

@app.middleware("http")
async def _no_stale_pages(request: Request, call_next):
    """Browseren skal altid tjekke for nye sider/filer (fx efter en opdatering af Pi'en)."""
    response = await call_next(request)
    response.headers.setdefault("Cache-Control", "no-cache")
    return response


bus = events.EventBus()  # hændelser til moduler og skærmens scan-besked
_capture = {"until": 0, "chip": None}
_cap_lock = threading.Lock()
_devtest = {"until": 0, "reads": []}  # udstyrstest: scanninger opsamles her og tæller ikke i løbet


# ------------------------------------------------------------------ chip-læsninger

def handle_read(chip, reader, ts=None):
    """Indgang for ALLE læsninger (primær, backup, manuel, simuleret)."""
    ts = ts or time.time()
    chip = readers.normalize(chip)
    # Rå tekstlog som ekstra sikkerhed ved siden af databasen
    with open(READ_LOG, "a") as f:
        f.write(f"{ts:.3f}\t{reader}\t{chip}\n")

    with _cap_lock:
        if _capture["until"] > ts:
            _capture.update(chip=chip, until=0)
            bus.publish({"type": "chip.registered", "reader": reader})
            return {"captured": chip}
        if _devtest["until"] > ts:
            _devtest["reads"] = (_devtest["reads"] + [{"reader": reader, "chip": chip, "ts": ts}])[-50:]
            return {"test": True}

    race = db.active_race()
    race_id = race["id"] if race else None
    p = db.q1("SELECT id FROM participants WHERE race_id=? AND chip=?", (race_id, chip)) if race else None
    window = 0 if (race or {}).get("test", {}) and race["test"].get("no_dedup") else db.setting("dedup_sec", 30)
    dup = window and db.q1("SELECT id FROM reads WHERE race_id IS ? AND chip=? AND duplicate=0 AND deleted=0 AND ABS(ts-?)<?",
                (race_id, chip, ts, window))
    rid = db.ex("INSERT INTO reads(race_id, chip, reader, ts, participant_id, duplicate) VALUES (?,?,?,?,?,?)",
                (race_id, chip, reader, ts, p["id"] if p else None, 1 if dup else 0))
    bus.read_stored(rid)  # vurderes i baggrunden – tidtagningen venter aldrig på moduler
    return {"id": rid, "participant_id": p["id"] if p else None, "duplicate": bool(dup)}


reader_mgr = readers.ReaderManager(handle_read)
module_mgr = modules.ModuleManager()
bus.subscribe(module_mgr.on_event)


@app.on_event("startup")
def _startup():
    db.conn()
    branding.import_local_once()
    if MIRROR:
        return
    sync_mgr.start()
    module_mgr.start()
    threading.Thread(target=sync_bootcode, daemon=True).start()
    events.Ticker(bus, standings, reader_mgr.status).start()
    paths = db.setting("readers", {}) or {}
    for role in readers.ROLES:  # miljøvariabler kan overstyre, fx i systemd-filen
        env = os.environ.get(f"RACETIMER_{role.upper()}")
        if env:
            paths[role] = env
    reader_mgr.configure(paths)


# ------------------------------------------------------------------ auth

def admin(request: Request):
    if MIRROR:
        raise HTTPException(403, "Dette er en spejlserver – administrér løbet på Pi'en")
    u = db.session_user(request.cookies.get("rt_session"))
    if not u:
        raise HTTPException(401, "Log ind")
    return u


def is_remote(request):
    """Kom forespørgslen ind via Cloudflare-tunnelen (fjernadgang fra internettet)?

    Headerne kan forfalskes på det lokale net – men så behandles forespørgslen bare strengere.
    """
    return bool(request.headers.get("cf-ray") or request.headers.get("cf-connecting-ip"))


def client_ip(request):
    """Afsenderens IP. Cloudflares header stoles kun på, når den kommer fra cloudflared på Pi'en selv –
    ellers kunne man på det lokale net omgå spærringen efter forkerte login ved at skifte header."""
    ip = request.client.host if request.client else "?"
    cf = request.headers.get("cf-connecting-ip")
    return cf if cf and ip in ("127.0.0.1", "::1") else ip


def is_https(request):
    return (request.url.scheme == "https" or request.headers.get("x-forwarded-proto") == "https"
            or "https" in (request.headers.get("cf-visitor") or ""))


_login_fails = {}  # ip -> (antal fejl, spærret til)


@app.post("/api/login")
async def login(request: Request, response: Response):
    if MIRROR:
        raise HTTPException(403, "Dette er en spejlserver – administrér løbet på Pi'en")
    ip = client_ip(request)
    fails, until = _login_fails.get(ip, (0, 0))
    if until > time.time():
        raise HTTPException(429, f"For mange forkerte forsøg – prøv igen om {int((until - time.time()) / 60) + 1} min")
    body = await request.json()
    user = body.get("username", "")
    if is_remote(request) and user == "admin" and db.default_password_in_use():
        raise HTTPException(403, "Førstegangskoden kan ikke bruges via fjernadgang – skift koden lokalt først")
    token = db.login(user, body.get("password", ""))
    if token and is_remote(request) and len(body.get("password", "")) < db.MIN_PW:
        db.ex("DELETE FROM sessions WHERE token=?", (token,))
        raise HTTPException(403, f"Koden er for kort til login via det offentlige link – skift den lokalt til mindst {db.MIN_PW} tegn")
    if not token:
        fails += 1
        _login_fails[ip] = (fails, time.time() + 15 * 60 if fails >= 5 else 0)
        await asyncio.sleep(1)  # ikke time.sleep: den ville fryse hele serveren (også resultatskærmen)
        raise HTTPException(401, "Forkert brugernavn eller kode")
    _login_fails.pop(ip, None)
    response.set_cookie("rt_session", token, httponly=True, samesite="strict", max_age=12 * 3600,
                        secure=is_https(request))
    return {"ok": True}


@app.post("/api/logout")
def logout(request: Request, response: Response):
    db.ex("DELETE FROM sessions WHERE token=?", (request.cookies.get("rt_session"),))
    response.delete_cookie("rt_session")
    return {"ok": True}


@app.get("/api/me")
def me(u=Depends(admin)):
    return {"id": u["id"], "username": u["username"], "default_password": db.default_password_in_use()}


@app.get("/api/setup")
def setup_status():
    """Offentlig: er opsætningsguiden gennemført? (Eksisterende installationer uden flaget regnes som færdige.)"""
    return {"done": bool(db.setting("setup_done", True)), "mirror": MIRROR}


@app.post("/api/setup/done")
def setup_done(u=Depends(admin)):
    if db.default_password_in_use():
        raise HTTPException(400, "Skift admin-koden først")
    db.set_setting("setup_done", True)
    return {"ok": True}


# ------------------------------------------------------------------ sider + live-data

@app.get("/")
def display_page():
    return FileResponse(os.path.join(STATIC, "display.html"))


@app.get("/admin")
def admin_page():
    return FileResponse(os.path.join(STATIC, "admin.html"))


def standings(race, now=None):
    now = now or time.time()
    parts = db.q("SELECT * FROM participants WHERE race_id=?", (race["id"],))
    reads = db.q("SELECT id, participant_id, ts, approved FROM reads WHERE race_id=? AND duplicate=0 AND deleted=0 "
                 "AND participant_id IS NOT NULL", (race["id"],))
    rows, info = engine.compute(race, parts, reads, now)
    s = signup.settings(race)
    if s["age_step"]:  # aldersklasser ud fra køn og fødselsår
        cls = {p["id"]: signup.age_class(p, race, s) for p in parts}
        for r in rows:
            r["class"] = cls.get(r["id"])
    # Afslut løbet automatisk når motoren siger det er afgjort
    if not MIRROR and race["status"] == "running" and info.get("phase") == "ended" and not race.get("finished_at"):
        fin = info.get("resolved_at") or info.get("end_at") or now
        db.ex("UPDATE races SET status='finished', finished_at=? WHERE id=?", (fin, race["id"]))
        bus.publish({"type": "race.finished", "race": race["name"]})
    return rows, info


sync_mgr = sync.SyncManager(standings)


def race_public(race):
    return {k: race.get(k) for k in ("id", "name", "type", "status", "scheduled_start", "started_at", "finished_at", "test")}


def on_pi_screen(request):
    """Kommer forespørgslen fra Pi'ens egen skærm (kiosk-browseren på localhost)?"""
    host = request.client.host if request.client else ""
    return host in ("127.0.0.1", "::1") and not is_remote(request)


@app.get("/api/live")
def live(request: Request):
    race = db.active_race()
    now = time.time()
    extra = {}
    if MIRROR:
        # brug Pi'ens ur, så nedtællinger matcher stævnet selv om urene afviger lidt
        now -= db.setting("mirror_clock_offset", 0)
        last = db.setting("mirror_last_push")
        extra["mirror"] = {"age": time.time() - last if last else None}
    extra["version"] = APP_VERSION
    extra["brand"] = branding.public()
    if not MIRROR and on_pi_screen(request):
        # admin-adresse, IP og WiFi kun på Pi'ens egen skærm – ikke for publikum på mobiler eller via fjernadgang
        extra["host"] = host_info()
        if db.first_code():  # ny installation: vis førstegangskoden på Pi'ens egen skærm
            extra["first_code"] = db.first_code()
    if not MIRROR:
        extra["clock"] = clock_info()
        extra["events"] = bus.recent()
    if not race:
        return {"race": None, "now": now, **extra}
    rows, info = standings(race, now)
    # kun de læsere, der er valgt – kører man med én læser, vises der ingen "backup"
    rd = [] if MIRROR else [{"role": r["role"], "connected": r["connected"]} for r in reader_mgr.status() if r["path"]]
    return {"race": race_public(race), "info": info, "rows": rows, "now": now, "readers": rd, **extra}


# ------------------------------------------------------------------ spejling

@app.post("/api/mirror/push")
async def mirror_push(request: Request):
    """Modtag data fra Pi'en (kun i spejl-tilstand, kræver delt nøgle)."""
    key = os.environ.get("RACETIMER_MIRROR_KEY", "")
    auth = request.headers.get("authorization", "")
    if not MIRROR or not key or not secrets.compare_digest(auth, "Bearer " + key):
        raise HTTPException(403, "Afvist")
    body = await request.body()
    if request.headers.get("content-encoding") == "gzip":
        body = gzip.decompress(body)
    b = json.loads(body)
    now = time.time()
    db.set_setting("mirror_clock_offset", now - b["pi_now"])
    db.set_setting("mirror_last_push", now)
    brand = b.get("branding")
    if brand and (brand.get("v") != db.setting("mirror_brand_v") or brand.get("images")):
        try:
            branding.apply_export(brand)
            db.set_setting("mirror_brand_v", brand.get("v"))
        except (ValueError, KeyError, TypeError):
            pass  # et ugyldigt billede må aldrig stoppe resultaterne
    race = b.get("race")
    if not race:
        db.ex("UPDATE races SET is_active=0")
        return {"ok": True}

    rid = race["id"]
    revs = db.setting("mirror_rev", {}) or {}
    known = revs.get(str(rid), 0)
    if b["since"] and b["since"] != known:
        return {"need_full": True}  # serveren mangler data (fx ny database) → Pi'en sender alt igen

    with db._lock:
        c = db.conn()
        c.execute("BEGIN")
        try:
            c.execute("INSERT INTO races(id, name, type, config, status, scheduled_start, started_at, finished_at, "
                      "created_at, signup) VALUES (?,?,?,?,?,?,?,?,?,?) ON CONFLICT(id) DO UPDATE SET name=excluded.name, "
                      "type=excluded.type, config=excluded.config, status=excluded.status, "
                      "scheduled_start=excluded.scheduled_start, started_at=excluded.started_at, "
                      "finished_at=excluded.finished_at, signup=excluded.signup",
                      (rid, race["name"], race["type"], json.dumps(race["config"]), race["status"],
                       race["scheduled_start"], race["started_at"], race["finished_at"], now,
                       json.dumps({"age_step": race.get("age_step") or 0, "age_from": race.get("age_from") or 35})))
            c.execute("UPDATE races SET is_active=(id=?)", (rid,))
            if b["since"] == 0:
                c.execute("DELETE FROM reads WHERE race_id=?", (rid,))
            if b.get("participants") is not None:
                c.execute("DELETE FROM participants WHERE race_id=?", (rid,))
                c.executemany("INSERT INTO participants(id, race_id, bib, name, sex, club, status, status_at, "
                              "start_offset, dogs, dog1, dog2, distance, birth_year) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                              [(p["id"], rid, p["bib"], p["name"], p["sex"], p["club"], p["status"], p["status_at"],
                                p["start_offset"], p.get("dogs"), p.get("dog1"), p.get("dog2"), p.get("distance"),
                                p.get("birth_year"))
                               for p in b["participants"]])
            c.executemany("INSERT OR REPLACE INTO reads(id, race_id, chip, reader, ts, participant_id, duplicate, "
                          "deleted, approved) VALUES (?,?,'',?,?,?,?,?,?)",
                          [(r[0], rid, r[3], r[2], r[1], r[4], r[5], r[6] if len(r) > 6 else 0)
                           for r in b.get("reads", [])])
            c.execute("COMMIT")
        except Exception:
            c.execute("ROLLBACK")
            raise
    revs[str(rid)] = b["rev"]
    db.set_setting("mirror_rev", revs)
    return {"ok": True}


# ------------------------------------------------------------------ synkronisering (Pi → webserver)

def _dest_public(d):
    st = sync_mgr.status.get(d["id"], {})
    out = {k: v for k, v in d.items() if k != "key"}
    out.update(key_set=bool(d.get("key")), last_ok=st.get("last_ok"), last_error=st.get("last_error"),
               last_event=st.get("last_event"), next_due=st.get("next_due") if d.get("interval_sec") else None,
               admin_announced=st.get("admin_announced"), admin_error=st.get("admin_error"),
               pending=sync_mgr.pending(d))
    return out


def _dest_fields(b, old=None):
    d = dict(old or {"id": secrets.token_hex(4)})
    typ = b.get("type", d.get("type"))
    if typ not in sync.TYPES:
        raise HTTPException(400, "Ukendt type")
    url = ((b["url"] if "url" in b else d.get("url")) or "").strip()
    schemes = ("ftps://", "ftp://") if typ == "ftp" else ("http://", "https://")
    if url and not url.startswith(schemes):
        raise HTTPException(400, "FTP-adressen skal starte med ftps:// eller ftp://" if typ == "ftp"
                            else "Adressen skal starte med https://")
    interval = float(b.get("interval_sec", d.get("interval_sec", 60)) or 0)
    if 0 < interval < 5:
        raise HTTPException(400, "Interval skal være mindst 5 sekunder (eller 0 = kun manuelt)")
    d.update(name=(b.get("name") or d.get("name") or "Webserver").strip(), type=typ, url=url,
             interval_sec=interval, on_finish=bool(b.get("on_finish", d.get("on_finish", True))),
             enabled=bool(b.get("enabled", d.get("enabled", True))),
             announce_admin=typ == "webhook" and bool(b.get("announce_admin", d.get("announce_admin", False))))
    if "username" in b:
        d["username"] = (b.get("username") or "").strip()
    if b.get("key"):
        d["key"] = b["key"] if typ == "ftp" else b["key"].strip()  # FTP-koder kan have mellemrum
    return d


@app.get("/api/destinations")
def list_destinations(u=Depends(admin)):
    return [_dest_public(d) for d in sync.destinations()]


@app.post("/api/destinations")
async def add_destination(request: Request, u=Depends(admin)):
    d = _dest_fields(await request.json())
    sync.save_destinations(sync.destinations() + [d])
    return {"id": d["id"]}


@app.put("/api/destinations/{did}")
async def update_destination(did: str, request: Request, u=Depends(admin)):
    dests = sync.destinations()
    i = next((i for i, d in enumerate(dests) if d["id"] == did), None)
    if i is None:
        _404()
    new = _dest_fields(await request.json(), dests[i])
    if (new["url"], new["type"]) != (dests[i]["url"], dests[i]["type"]):
        sync_mgr.reset(did)  # ny server → send alt forfra
    dests[i] = new
    sync.save_destinations(dests)
    return {"ok": True}


@app.delete("/api/destinations/{did}")
def delete_destination(did: str, u=Depends(admin)):
    sync.save_destinations([d for d in sync.destinations() if d["id"] != did])
    sync_mgr.reset(did)
    return {"ok": True}


@app.post("/api/destinations/{did}/send")
def send_destination(did: str, u=Depends(admin)):
    """Manuel afsendelse. Er løbet afsluttet, sendes det som endelige resultater."""
    d = next((d for d in sync.destinations() if d["id"] == did), None) or _404()
    if not d.get("url"):
        raise HTTPException(400, "Destinationen har ingen adresse")
    race = db.active_race()
    final = bool(race and race["status"] == "finished")
    if not sync_mgr.send(d, final=final, manual=True):
        raise HTTPException(502, sync_mgr.status.get(did, {}).get("last_error") or "Afsendelse fejlede")
    return {"ok": True, "final": final}

# ------------------------------------------------------------------ løb

@app.get("/api/types")
def types(u=Depends(admin)):
    return engine.RACE_TYPES


# ------------------------------------------------------------------ skabeloner (egne løbsopsætninger)

def _clean_template(b):
    """En skabelon: løbstype, indstillinger (inkl. ruter og distancer) og tæller – ingen deltagere eller tider."""
    typ = b.get("type")
    if typ not in engine.RACE_TYPES:
        raise HTTPException(400, "Ukendt løbstype i skabelonen")
    name = str(b.get("name") or "").strip()[:80]
    if not name:
        raise HTTPException(400, "Skabelonen skal have et navn")
    cfg = b.get("config") if isinstance(b.get("config"), dict) else {}
    return {"name": name, "description": str(b.get("description") or "").strip()[:300], "type": typ,
            "timing": "manual" if b.get("timing") == "manual" else "armband",
            "config": engine.clean_distances(typ, engine.merged_config(typ, cfg))}


def _templates():
    return db.setting("templates", []) or []


@app.get("/api/templates")
def templates(u=Depends(admin)):
    return {"examples": engine.EXAMPLE_TEMPLATES, "own": _templates()}


@app.post("/api/templates")
async def save_template(request: Request, u=Depends(admin)):
    """Gem en skabelon – fra et løb ({"race_id", "name"}) eller fra en fil/formular (type, config, …)."""
    b = await request.json()
    if b.get("race_id"):
        race = db.race(int(b["race_id"])) or _404()
        b = {**b, "type": race["type"], "config": race["config"], "timing": race.get("timing")}
    t = _clean_template(b)
    own = _templates()
    old = next((x for x in own if x["name"].lower() == t["name"].lower()), None)
    if old:  # samme navn → opdatér (fx næste års udgave)
        old.update(t, updated_at=time.time())
    else:
        t.update(id=secrets.token_hex(4), created_at=time.time())
        own.append(t)
    db.set_setting("templates", own)
    return old or t


@app.delete("/api/templates/{tid}")
def delete_template(tid: str, u=Depends(admin)):
    own = _templates()
    if not any(x["id"] == tid for x in own):
        _404()
    db.set_setting("templates", [x for x in own if x["id"] != tid])
    return {"ok": True}


@app.get("/api/templates/{tid}/export")
def export_template(tid: str, u=Depends(admin)):
    t = next((x for x in _templates() + engine.EXAMPLE_TEMPLATES if x["id"] == tid), None) or _404()
    out = {"racetimer_template": 1, **{k: t[k] for k in ("name", "description", "type", "timing", "config")}}
    fn = "".join(c if c.isalnum() else "-" for c in t["name"].lower()).strip("-") or "skabelon"
    return Response(json.dumps(out, ensure_ascii=False, indent=2), media_type="application/json",
                    headers={"Content-Disposition": f'attachment; filename="skabelon-{fn}.json"'})


@app.get("/api/races")
def races(u=Depends(admin)):
    rows = db.q("SELECT r.*, (SELECT COUNT(*) FROM participants p WHERE p.race_id=r.id) AS n "
                "FROM races r ORDER BY r.created_at DESC")
    for r in rows:
        r["config"] = json.loads(r["config"])
        r["test"] = json.loads(r["test"]) if r.get("test") else None
    return rows


@app.get("/api/races/{rid}")
def race_detail(rid: int, u=Depends(admin)):
    race = db.race(rid) or _404()
    rows, info = standings(race)
    race["effective_config"] = engine.merged_config(race["type"], race["config"])
    return {"race": race, "rows": rows, "info": info}


def _404():
    raise HTTPException(404, "Findes ikke")


@app.post("/api/races")
async def create_race(request: Request, u=Depends(admin)):
    b = await request.json()
    if b.get("type") not in engine.RACE_TYPES:
        raise HTTPException(400, "Ukendt løbstype")
    cfg = engine.clean_distances(b["type"], engine.merged_config(b["type"], b.get("config")))
    test = _test_opts(b.get("test"))
    timing = "manual" if b.get("timing") == "manual" else "armband"
    rid = db.ex("INSERT INTO races(name, type, config, scheduled_start, created_at, test, timing) VALUES (?,?,?,?,?,?,?)",
                (b.get("name") or "Nyt løb", b["type"], json.dumps(cfg), b.get("scheduled_start"), time.time(),
                 json.dumps(test) if test else None, timing))
    if b.get("activate") or not db.active_race():
        _activate(rid)
    return {"id": rid}


@app.put("/api/races/{rid}")
async def update_race(rid: int, request: Request, u=Depends(admin)):
    race = db.race(rid) or _404()
    b = await request.json()
    cfg = engine.clean_distances(race["type"], engine.merged_config(race["type"], {**race["config"], **(b.get("config") or {})}))
    db.ex("UPDATE races SET name=?, config=?, scheduled_start=? WHERE id=?",
          (b.get("name", race["name"]), json.dumps(cfg), b.get("scheduled_start", race["scheduled_start"]), rid))
    if race.get("test") and b.get("test"):  # et testløb forbliver et testløb – kun indstillingerne kan ændres
        db.ex("UPDATE races SET test=? WHERE id=?", (json.dumps(_test_opts(b["test"]) or {"speed": 1}), rid))
    return {"ok": True}


def _test_opts(t):
    """Testløbets indstillinger (None = rigtigt løb)."""
    if not t:
        return None
    speed = int(t.get("speed") or 1)
    return {"speed": speed if speed in engine.TEST_SPEEDS else 1, "no_minimum": bool(t.get("no_minimum")),
            "no_checkin": bool(t.get("no_checkin")), "no_dedup": bool(t.get("no_dedup"))}


@app.delete("/api/races/{rid}")
def delete_race(rid: int, u=Depends(admin)):
    db.ex("DELETE FROM races WHERE id=?", (rid,))
    return {"ok": True}


def _activate(rid):
    db.ex("UPDATE races SET is_active = (id=?)", (rid,))
    sync_mgr.poke_state()


@app.post("/api/races/deactivate")
def deactivate(u=Depends(admin)):
    """Intet aktivt løb → resultatskærmen viser splash-billedet."""
    db.ex("UPDATE races SET is_active=0")
    sync_mgr.poke_state()
    return {"ok": True}


@app.post("/api/races/{rid}/activate")
def activate(rid: int, u=Depends(admin)):
    db.ex("UPDATE races SET archived_at=NULL WHERE id=?", (rid,))  # det, der vises, hører ikke til i arkivet
    _activate(rid)
    return {"ok": True}


@app.post("/api/races/{rid}/start")
async def start_race(rid: int, request: Request, u=Depends(admin)):
    b = await request.json() if await request.body() else {}
    at = b.get("at") or time.time()
    db.ex("UPDATE races SET status='running', started_at=?, finished_at=NULL WHERE id=?", (at, rid))
    bus.publish({"type": "race.started", "race": (db.race(rid) or {}).get("name")})
    return {"ok": True, "started_at": at}


@app.post("/api/races/{rid}/finish")
def finish_race(rid: int, u=Depends(admin)):
    db.ex("UPDATE races SET status='finished', finished_at=? WHERE id=?", (time.time(), rid))
    bus.publish({"type": "race.finished", "race": (db.race(rid) or {}).get("name")})
    return {"ok": True}


@app.post("/api/races/{rid}/resume")
def resume_race(rid: int, u=Depends(admin)):
    db.ex("UPDATE races SET status='running', finished_at=NULL WHERE id=?", (rid,))
    return {"ok": True}


# ------------------------------------------------------------------ tæller: armbånd eller manuel (tryk på felter)

TAP_DOUBLE_SEC = 3  # et nyt tryk på samme løber inden for 3 sek. er et dobbelttryk


@app.post("/api/races/{rid}/timing")
async def set_timing(rid: int, request: Request, u=Depends(admin)):
    """Skift tæller – også midt i et løb. Alle registreringer bevares (tryk og scanninger tæller ens)."""
    db.race(rid) or _404()
    mode = (await request.json()).get("timing")
    if mode not in ("armband", "manual"):
        raise HTTPException(400, "Vælg armbånd eller manuel")
    db.ex("UPDATE races SET timing=? WHERE id=?", (mode, rid))
    return {"ok": True, "timing": mode}


@app.post("/api/races/{rid}/tap")
async def tap(rid: int, request: Request, u=Depends(admin)):
    """Manuel tidtagning: race director trykker på løberens felt. Tæller som en chip-scanning
    (minimumstid, tidsvinduer og indtjekning gælder) – og giver samme scan-besked og lyd."""
    race = db.race(rid) or _404()
    pid = int((await request.json()).get("participant_id") or 0)
    p = db.q1("SELECT id, chip FROM participants WHERE id=? AND race_id=?", (pid, rid)) or _404()
    ts = time.time()
    dup = db.q1("SELECT id FROM reads WHERE race_id=? AND participant_id=? AND duplicate=0 AND deleted=0 AND ABS(ts-?)<?",
                (rid, pid, ts, TAP_DOUBLE_SEC))
    if (race.get("test") or {}).get("no_dedup"):
        dup = None
    read_id = db.ex("INSERT INTO reads(race_id, chip, reader, ts, participant_id, duplicate) VALUES (?,?,?,?,?,?)",
                    (rid, p["chip"] or f"TRYK-{pid}", "tap", ts, pid, 1 if dup else 0))
    try:
        ev = outcome.evaluate(read_id)
    except Exception:  # noqa: BLE001 – vurderingen må aldrig stoppe registreringen
        ev = None
    if ev and race["is_active"]:  # scan-besked og lyd kun for løbet på skærmen
        ev = bus.publish(ev)
    return {"id": read_id, "event": ev}


@app.post("/api/races/{rid}/tap/undo")
def tap_undo(rid: int, u=Depends(admin)):
    """Fortryd det seneste tryk, der talte (soft delete – kan gendannes i loggen over læsninger)."""
    last = db.q1("SELECT r.id, r.ts, p.name, p.bib FROM reads r LEFT JOIN participants p ON p.id=r.participant_id "
                 "WHERE r.race_id=? AND r.reader='tap' AND r.deleted=0 AND r.duplicate=0 ORDER BY r.ts DESC LIMIT 1", (rid,))
    if not last:
        raise HTTPException(404, "Intet tryk at fortryde")
    db.ex("UPDATE reads SET deleted=1 WHERE id=?", (last["id"],))
    return {"ok": True, **last}


# ------------------------------------------------------------------ tilmelding på hjemmesiden

def _signup_dests():
    return [d for d in sync.destinations() if d.get("type") == "webhook" and d.get("enabled") and d.get("url")]


@app.get("/api/races/{rid}/signup")
def get_signup(rid: int, u=Depends(admin)):
    race = db.race(rid) or _404()
    s = signup.settings(race)
    dests = _signup_dests()
    online = db.q1("SELECT COUNT(*) n FROM participants WHERE race_id=? AND signup_id IS NOT NULL", (rid,))["n"]
    return {**s, "is_open": signup.is_open(race, s), "online": online, "test": bool(race.get("test")),
            "purge_at": signup.purge_at(race, s),
            "forms": [re.sub(r"modtag\.php(\?.*)?$", "tilmeld.php", d["url"]) + f"?lob={rid}" for d in dests],
            "errors": [sync_mgr.status.get(d["id"], {}).get("signup_error") for d in dests if
                       sync_mgr.status.get(d["id"], {}).get("signup_error")]}


@app.put("/api/races/{rid}/signup")
async def put_signup(rid: int, request: Request, u=Depends(admin)):
    race = db.race(rid) or _404()
    s = signup.clean(await request.json(), signup.settings(race))
    if s["open"] and race.get("test"):
        raise HTTPException(400, "Testløb kan ikke åbnes for tilmelding")
    signup.save(rid, s)
    sync_mgr.poke_state()  # hjemmesiden får besked med det samme
    return get_signup(rid, u)


@app.post("/api/races/{rid}/signup/fetch")
def fetch_signup(rid: int, u=Depends(admin)):
    db.race(rid) or _404()
    dests = _signup_dests()
    if not dests:
        raise HTTPException(400, "Opret først en destination af typen “Egen hjemmeside (JSON)” under ⚙ Opsætning → Synkronisering")
    added = 0
    for d in dests:
        try:
            added += sync_mgr.fetch_signups(d, force=True) or 0
        except RuntimeError as e:
            raise HTTPException(502, f"Kunne ikke hente fra {d.get('name') or d['url']}: {e}") from None
    return {"added": added}


@app.post("/api/races/{rid}/archive")
def archive_race(rid: int, u=Depends(admin)):
    """Flyt et afsluttet løb til arkivet. Intet slettes, og det kan gendannes."""
    race = db.race(rid) or _404()
    if race["status"] != "finished":
        raise HTTPException(400, "Kun afsluttede løb kan arkiveres")
    db.ex("UPDATE races SET archived_at=?, is_active=0 WHERE id=?", (time.time(), rid))
    sync_mgr.poke_state()  # hjemmesiden skal skjule løbet med det samme
    return {"ok": True}


@app.post("/api/races/{rid}/unarchive")
def unarchive_race(rid: int, u=Depends(admin)):
    db.race(rid) or _404()
    db.ex("UPDATE races SET archived_at=NULL WHERE id=?", (rid,))
    sync_mgr.poke_state()
    return {"ok": True}


@app.post("/api/races/{rid}/reset")
def reset_race(rid: int, u=Depends(admin)):
    """Nulstil start. Læsninger markeres som slettede (kan gendannes) – intet slettes permanent."""
    db.ex("UPDATE races SET status='planned', started_at=NULL, finished_at=NULL WHERE id=?", (rid,))
    db.ex("UPDATE reads SET deleted=1 WHERE race_id=?", (rid,))
    db.ex("UPDATE participants SET status=NULL, status_at=NULL WHERE race_id=?", (rid,))
    return {"ok": True}


@app.get("/api/races/{rid}/results.csv")
def results_csv(rid: int, u=Depends(admin)):
    race = db.race(rid) or _404()
    rows, info = standings(race)
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    cfg = engine.merged_config(race["type"], race["config"])
    cc, multi = bool(cfg.get("canicross")), bool(engine.distances(race["type"], cfg))
    ages = bool(signup.settings(race)["age_step"])
    w.writerow((["Distance"] if multi else []) + ["Placering", "Startnr", "Navn", "Køn"] + (["Klasse"] if ages else []) + ["Klub"] +
               (["Hunde"] if cc else []) +
               ["Status", "Omgange", "Distance km", "Tid (s)", "Sidste omgang (s)", "Bedste omgang (s)"])
    for r in rows:
        w.writerow(([r.get("distance_name")] if multi else []) + [r["place"], r["bib"], r["name"], r["sex"]] +
                   ([r.get("class") or ""] if ages else []) + [r["club"]] +
                   ([" & ".join(r["dogs"])] if cc else []) +
                   [r["status_label"], r["laps"], r["distance_km"], r["total_time"] and round(r["total_time"], 1),
                    r["last_lap"] and round(r["last_lap"], 1), r["best_lap"] and round(r["best_lap"], 1)])
    return Response("﻿" + buf.getvalue(), media_type="text/csv",
                    headers={"Content-Disposition": f'attachment; filename="resultater-{rid}.csv"'})


# ------------------------------------------------------------------ deltagere

SEX = {"m": "Mand", "male": "Mand", "mand": "Mand", "herre": "Mand", "h": "Mand",
       "k": "Kvinde", "f": "Kvinde", "female": "Kvinde", "kvinde": "Kvinde", "dame": "Kvinde", "d": "Kvinde"}


def norm_sex(s):
    s = (s or "").strip()
    return SEX.get(s.lower(), s)


def _participant_fields(b):
    bib = b.get("bib")
    dogs = str(b.get("dogs") or "").strip()
    return {
        "bib": int(bib) if str(bib or "").strip().isdigit() else None,
        "name": (b.get("name") or "").strip(),
        "sex": norm_sex(b.get("sex")),
        "club": (b.get("club") or "").strip(),
        "chip": readers.normalize(b.get("chip")) or None,
        "dogs": int(dogs) if dogs in ("1", "2") else None,
        "dog1": (b.get("dog1") or "").strip() or None,
        "dog2": (b.get("dog2") or "").strip() or None,
        "distance": str(b.get("distance") or "").strip() or None,
        "birth_year": _birth_year(b.get("birth_year")),
        "email": (str(b.get("email") or "").strip()[:200] or None),
    }


def _birth_year(v):
    v = str(v or "").strip().removesuffix(".0")
    return int(v) if v.isdigit() and 1900 <= int(v) <= time.localtime().tm_year else None


def _distance_id(rid, value):
    """Distance ud fra id eller navn (fx fra Excel: "10 km"). None = løbets første distance."""
    race = db.race(rid)
    ds = engine.distances(race["type"], engine.merged_config(race["type"], race["config"])) if race else []
    v = str(value or "").strip().lower()
    if not v or not ds:
        return None
    hit = next((d["id"] for d in ds if v in (d["id"].lower(), d["name"].lower())), None)
    if not hit:
        raise HTTPException(400, f"Ukendt distance: {value} (vælg mellem {', '.join(d['name'] for d in ds)})")
    return hit


def _check_dogs(rid, f):
    """Canicross: hver deltager skal have 1 eller 2 hunde med navn."""
    race = db.race(rid)
    if not race or not engine.merged_config(race["type"], race["config"]).get("canicross"):
        return
    if f["dogs"] is None:  # fx fra Excel uden antal: tæl de udfyldte navne
        f["dogs"] = len([d for d in (f["dog1"], f["dog2"]) if d]) or None
    if f["dogs"] not in (1, 2):
        raise HTTPException(400, f"{f['name'] or 'Deltageren'}: vælg 1 eller 2 hunde (canicross)")
    if not f["dog1"]:
        raise HTTPException(400, f"{f['name']}: skriv navnet på hunden")
    if f["dogs"] == 2 and not f["dog2"]:
        raise HTTPException(400, f"{f['name']}: skriv navnet på hund 2")
    if f["dogs"] == 1:
        f["dog2"] = None


def _check_chip(rid, chip, pid=None):
    if chip:
        other = db.q1("SELECT name FROM participants WHERE race_id=? AND chip=? AND id IS NOT ?", (rid, chip, pid))
        if other:
            raise HTTPException(409, f"Chip {chip} er allerede tildelt {other['name']}")


@app.get("/api/races/{rid}/participants")
def participants(rid: int, u=Depends(admin)):
    return db.q("SELECT * FROM participants WHERE race_id=? ORDER BY bib IS NULL, bib, name", (rid,))


@app.post("/api/races/{rid}/participants")
async def add_participant(rid: int, request: Request, u=Depends(admin)):
    f = _participant_fields(await request.json())
    if not f["name"]:
        raise HTTPException(400, "Navn mangler")
    _check_chip(rid, f["chip"])
    _check_dogs(rid, f)
    f["distance"] = _distance_id(rid, f["distance"])
    pid = db.ex("INSERT INTO participants(race_id, bib, name, sex, club, chip, dogs, dog1, dog2, distance, birth_year, email) "
                "VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
                (rid, f["bib"], f["name"], f["sex"], f["club"], f["chip"], f["dogs"], f["dog1"], f["dog2"], f["distance"],
                 f["birth_year"], f["email"]))
    _relink(rid)
    _bibs_changed(rid)
    return {"id": pid}


@app.put("/api/participants/{pid}")
async def update_participant(pid: int, request: Request, u=Depends(admin)):
    p = db.q1("SELECT * FROM participants WHERE id=?", (pid,)) or _404()
    b = await request.json()
    f = _participant_fields({**p, **b})
    _check_chip(p["race_id"], f["chip"], pid)
    if any(k in b for k in ("dogs", "dog1", "dog2")):  # ellers kan man ikke rette fx navnet på en deltager uden hunde
        _check_dogs(p["race_id"], f)
    so = b.get("start_offset", p["start_offset"])
    f["distance"] = _distance_id(p["race_id"], f["distance"])
    db.ex("UPDATE participants SET bib=?, name=?, sex=?, club=?, chip=?, start_offset=?, dogs=?, dog1=?, dog2=?, distance=?, "
          "birth_year=?, email=? WHERE id=?",
          (f["bib"], f["name"], f["sex"], f["club"], f["chip"], None if so in ("", None) else float(so),
           f["dogs"], f["dog1"], f["dog2"], f["distance"], f["birth_year"], f["email"], pid))
    _relink(p["race_id"])
    if f["bib"] != p["bib"]:
        _bibs_changed(p["race_id"])
    return {"ok": True}


@app.delete("/api/participants/{pid}")
def delete_participant(pid: int, u=Depends(admin)):
    db.ex("UPDATE reads SET participant_id=NULL WHERE participant_id=?", (pid,))
    db.ex("DELETE FROM participants WHERE id=?", (pid,))
    return {"ok": True}


@app.post("/api/participants/{pid}/status")
async def set_status(pid: int, request: Request, u=Depends(admin)):
    st = (await request.json()).get("status") or None
    if st not in (None, "RTC", "DSQ", "DNS"):
        raise HTTPException(400, "Ugyldig status")
    db.ex("UPDATE participants SET status=?, status_at=? WHERE id=?", (st, time.time() if st else None, pid))
    return {"ok": True}


def _bibs_changed(rid):
    """Startnumre på Pi'en ændret – åben tilmelding: hjemmesiden skal kende de optagne numre med det samme."""
    race = db.race(rid)
    if race and signup.is_open(race, signup.settings(race)):
        sync_mgr.poke_state()


def _relink(rid):
    """Knyt læsninger til deltagere ud fra chip – så en chip tildelt efter start stadig tæller."""
    db.ex("UPDATE reads SET participant_id=(SELECT id FROM participants p WHERE p.race_id=reads.race_id "
          "AND p.chip=reads.chip) WHERE race_id=? AND reader!='manual'", (rid,))


HEADERS = {
    "name": ("navn", "name", "fulde navn", "deltager"),
    "first": ("fornavn", "first name", "firstname"),
    "last": ("efternavn", "last name", "lastname", "surname"),
    "sex": ("køn", "kon", "sex", "gender", "m/k"),
    "bib": ("startnr", "startnummer", "bib", "nr", "nummer", "start no"),
    "chip": ("chip", "rfid", "tag", "chip id", "chipnr"),
    "club": ("klub", "club", "team", "forening"),
    "dogs": ("hunde", "antal hunde", "dogs"),
    "dog1": ("hund 1", "hund1", "hund", "hundenavn", "dog 1", "dog1", "dog"),
    "dog2": ("hund 2", "hund2", "dog 2", "dog2"),
    "distance": ("distance", "distancer", "løb", "rute"),
    "birth_year": ("fødselsår", "født", "årgang", "birth year", "year of birth"),
    "email": ("e-mail", "email", "mail"),
}


@app.post("/api/races/{rid}/import")
async def import_participants(rid: int, file: UploadFile = File(...), u=Depends(admin)):
    data = await file.read()
    if file.filename.lower().endswith((".xlsx", ".xlsm")):
        import openpyxl
        ws = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        table = [[("" if c is None else str(c)) for c in row] for row in ws.iter_rows(values_only=True)]
    else:
        text = data.decode("utf-8-sig", "ignore")
        table = list(csv.reader(io.StringIO(text), delimiter=";" if text.count(";") > text.count(",") else ","))
    if not table:
        raise HTTPException(400, "Tom fil")
    head = [h.strip().lower() for h in table[0]]
    col = {k: next((i for i, h in enumerate(head) if h in names), None) for k, names in HEADERS.items()}
    if col["name"] is None and col["first"] is None:
        raise HTTPException(400, "Kolonnen 'Navn' (eller 'Fornavn'/'Efternavn') blev ikke fundet")

    added = updated = 0
    errors = []
    for n, row in enumerate(table[1:], start=2):
        get = lambda k: row[col[k]].strip() if col[k] is not None and col[k] < len(row) else ""  # noqa: E731
        name = get("name") or f"{get('first')} {get('last')}".strip()
        if not name:
            continue
        if get("bib").endswith(".0"):  # Excel gemmer tal som float
            row[col["bib"]] = get("bib")[:-2]
        if get("chip").endswith(".0"):
            row[col["chip"]] = get("chip")[:-2]
        f = _participant_fields({"name": name, "sex": get("sex"), "bib": get("bib"), "club": get("club"),
                                 "chip": get("chip"), "dogs": get("dogs").replace(".0", ""),
                                 "dog1": get("dog1"), "dog2": get("dog2"), "distance": get("distance"),
                                 "birth_year": get("birth_year"), "email": get("email")})
        existing = None
        if f["bib"] is not None:
            existing = db.q1("SELECT id FROM participants WHERE race_id=? AND bib=?", (rid, f["bib"]))
        try:
            _check_chip(rid, f["chip"], existing["id"] if existing else None)
            _check_dogs(rid, f)
            f["distance"] = _distance_id(rid, f["distance"])
        except HTTPException as e:
            errors.append(f"Række {n}: {e.detail}")
            continue
        if existing:
            db.ex("UPDATE participants SET name=?, sex=?, club=?, chip=COALESCE(?, chip), dogs=COALESCE(?, dogs), "
                  "dog1=COALESCE(?, dog1), dog2=COALESCE(?, dog2), distance=COALESCE(?, distance), "
                  "birth_year=COALESCE(?, birth_year), email=COALESCE(?, email) WHERE id=?",
                  (f["name"], f["sex"], f["club"], f["chip"], f["dogs"], f["dog1"], f["dog2"], f["distance"],
                   f["birth_year"], f["email"], existing["id"]))
            updated += 1
        else:
            db.ex("INSERT INTO participants(race_id, bib, name, sex, club, chip, dogs, dog1, dog2, distance, birth_year, email) "
                  "VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
                  (rid, f["bib"], f["name"], f["sex"], f["club"], f["chip"], f["dogs"], f["dog1"], f["dog2"], f["distance"],
                   f["birth_year"], f["email"]))
            added += 1
    _relink(rid)
    _bibs_changed(rid)
    return {"added": added, "updated": updated, "errors": errors}


@app.get("/api/import-template.xlsx")
def import_template(u=Depends(admin)):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Deltagere"
    ws.append(["Startnr", "Navn", "Køn", "Fødselsår", "Klub", "E-mail", "Chip", "Distance", "Hunde", "Hund 1", "Hund 2"])
    ws.append([1, "Kenneth Risom", "Mand", 1978, "Aalborg Løbeklub", "", "", "", 1, "Bella", ""])
    ws.append([2, "Mette Hansen", "Kvinde", 1990, "", "", "", "", 2, "Max", "Luna"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="deltagere-skabelon.xlsx"'})


# ------------------------------------------------------------------ chip-registrering (capture)

@app.post("/api/capture/start")
def capture_start(u=Depends(admin)):
    with _cap_lock:
        _capture.update(until=time.time() + 30, chip=None)
    return {"ok": True}


@app.post("/api/capture/cancel")
def capture_cancel(u=Depends(admin)):
    with _cap_lock:
        _capture.update(until=0, chip=None)
    return {"ok": True}


# ------------------------------------------------------------------ udstyrstest

@app.post("/api/devtest/start")
def devtest_start(u=Depends(admin)):
    """Lyt efter scanninger i 2 min. De gemmes ikke som læsninger og tæller ikke i noget løb."""
    with _cap_lock:
        _devtest.update(until=time.time() + 120, reads=[])
    return {"ok": True, "now": time.time()}


@app.get("/api/devtest")
def devtest_poll(u=Depends(admin)):
    with _cap_lock:
        active = _devtest["until"] > time.time()
        reads = list(_devtest["reads"])
    return {"active": active, "reads": reads, "now": time.time(), "readers": reader_mgr.status(),
            "last": db.setting("devtest_last")}


@app.post("/api/devtest/stop")
def devtest_stop(u=Depends(admin)):
    with _cap_lock:
        _devtest.update(until=0)
    return {"ok": True}


@app.post("/api/devtest/screen")
def devtest_screen(u=Depends(admin)):
    """Vis en test-besked på alle åbne resultatskærme."""
    bus.publish({"type": "test.screen"})
    return {"ok": True}


@app.post("/api/devtest/result")
async def devtest_result(request: Request, u=Depends(admin)):
    """Gem resultatet af den seneste udstyrstest (vises næste gang fanen åbnes)."""
    b = await request.json()
    items = [{"name": str(i.get("name", ""))[:80], "ok": i.get("ok"), "note": str(i.get("note") or "")[:160]}
             for i in (b.get("items") or [])[:20]]
    last = {"ts": time.time(), "by": u["username"], "items": items}
    db.set_setting("devtest_last", last)
    return last


@app.get("/api/capture")
def capture_poll(u=Depends(admin)):
    with _cap_lock:
        chip = _capture["chip"]
        waiting = _capture["until"] > time.time()
        if chip:
            _capture["chip"] = None
    owner = None
    race = db.active_race()
    if chip and race:
        owner = db.q1("SELECT id, name, bib FROM participants WHERE race_id=? AND chip=?", (race["id"], chip))
    return {"chip": chip, "waiting": waiting, "owner": owner}


# ------------------------------------------------------------------ læsninger

@app.get("/api/races/{rid}/reads")
def list_reads(rid: int, limit: int = 200, u=Depends(admin)):
    return db.q("SELECT r.*, p.name, p.bib FROM reads r LEFT JOIN participants p ON p.id=r.participant_id "
                "WHERE r.race_id=? ORDER BY r.ts DESC LIMIT ?", (rid, limit))


@app.post("/api/races/{rid}/reads")
async def manual_read(rid: int, request: Request, u=Depends(admin)):
    """Manuel tid for en løber (fx hvis RFID fejler).

    Tiden angives som ts (klokkeslæt, epoch) eller elapsed (sek. løbstid fra løbets/løberens start).
    kind: "split" = omgang/mellemtid · "final" = sluttid (erstatter tidligere mållæsninger)
          "start" = starttid (tidskørsel med chipstart) · "checkin" = manuel indtjekning
    """
    b = await request.json()
    race = db.race(rid) or _404()
    p = db.q1("SELECT * FROM participants WHERE id=? AND race_id=?", (b.get("participant_id"), rid)) or _404()
    kind = b.get("kind") or "split"
    if kind not in ("split", "final", "start", "checkin"):
        raise HTTPException(400, "Ukendt type")
    row = next((r for r in standings(race)[0] if r["id"] == p["id"]), {})
    if kind == "checkin":
        # manuel indtjekning: en læsning lige før start (eller nu, hvis løbet ikke er startet)
        ts = race["started_at"] - 1 if race["started_at"] else time.time()
        if row.get("status") not in ("nocheckin", "waiting"):
            raise HTTPException(409, f"{p['name']} er allerede tjekket ind")
        db.ex("INSERT INTO reads(race_id, chip, reader, ts, participant_id, approved) VALUES (?,?,?,?,?,1)",
              (rid, p["chip"] or f"MANUEL-{p['id']}", "manual", ts, p["id"]))
        return {"ok": True, "checkin": True, "name": p["name"], "ts": ts}
    if b.get("elapsed") is not None:
        base = row.get("start_ts") or race["started_at"]
        if not base:
            raise HTTPException(400, "Løbet er ikke startet – angiv et klokkeslæt i stedet")
        ts = base + float(b["elapsed"])
    else:
        ts = float(b.get("ts") or time.time())
    if race["started_at"] and ts < race["started_at"] - 1 and kind != "start":
        raise HTTPException(400, "Tiden ligger før løbets start")
    if ts > time.time() + 60:
        raise HTTPException(400, "Tiden ligger i fremtiden")
    if not race["started_at"] and kind != "start":
        raise HTTPException(400, "Løbet er ikke startet endnu – start det først")
    dup = db.q1("SELECT ts FROM reads WHERE race_id=? AND participant_id=? AND reader='manual' AND deleted=0 "
                "AND ABS(ts-?)<10", (rid, p["id"], ts))
    if dup:  # beskytter mod dobbeltklik
        raise HTTPException(409, f"{p['name']} har allerede en manuel tid {abs(ts - dup['ts']):.0f} sek. fra denne")

    if kind == "final":
        # sluttiden skal være den, der tæller: andre læsninger efter løberens start ignoreres
        # (soft delete – kan gendannes i loggen over læsninger)
        cfg = engine.effective_config(race)
        after = (row.get("start_ts") or race["started_at"] or 0) + cfg.get("min_time_min", 0) * 60
        if race["type"] in ("p2p", "timetrial"):
            db.ex("UPDATE reads SET deleted=1 WHERE race_id=? AND participant_id=? AND ts>=? AND deleted=0",
                  (rid, p["id"], after))
    elif kind == "start":
        # starttid ved chipstart: tidligere læsninger må ikke "stjæle" starten
        db.ex("UPDATE reads SET deleted=1 WHERE race_id=? AND participant_id=? AND ts<? AND deleted=0",
              (rid, p["id"], ts))

    # manuel registrering er officialens egen afgørelse → tæller også i backyards flagzone
    rid_ = db.ex("INSERT INTO reads(race_id, chip, reader, ts, participant_id, approved) VALUES (?,?,?,?,?,1)",
                 (rid, p["chip"] or f"MANUEL-{p['id']}", "manual", ts, p["id"]))
    after = next((r for r in standings(race)[0] if r["id"] == p["id"]), {})
    counted = after.get("laps", 0) > row.get("laps", 0) or after.get("total_time") != row.get("total_time")
    return {"id": rid_, "ts": ts, "counted": counted, "laps": after.get("laps"), "total_time": after.get("total_time"),
            "name": p["name"]}


@app.post("/api/reads/{read_id}/approve")
def approve_read(read_id: int, u=Depends(admin)):
    """Official godkender/fortryder en læsning fra flagzonen som gennemført runde."""
    db.ex("UPDATE reads SET approved = 1 - approved WHERE id=?", (read_id,))
    return {"ok": True}


@app.post("/api/reads/{read_id}/toggle")
def toggle_read(read_id: int, u=Depends(admin)):
    """Ignorér/gendan en læsning (soft delete – data bevares altid)."""
    db.ex("UPDATE reads SET deleted = 1 - deleted WHERE id=?", (read_id,))
    return {"ok": True}


@app.post("/api/sim-read")
async def sim_read(request: Request, u=Depends(admin)):
    b = await request.json()
    return handle_read(b.get("chip", ""), b.get("reader") or "sim")


# ------------------------------------------------------------------ læsere + indstillinger

@app.get("/api/readers")
def reader_status(u=Depends(admin)):
    return {"status": reader_mgr.status(), "devices": readers.list_devices(), "ambiguous": readers.ambiguous_paths(),
            "config": db.setting("readers", {}), "dedup_sec": db.setting("dedup_sec", 30),
            "evdev": readers.evdev is not None}


@app.put("/api/readers")
async def set_readers(request: Request, u=Depends(admin)):
    b = await request.json()
    paths = {r: (b.get(r) or None) for r in readers.ROLES}
    if paths["primary"] and paths["backup"] and \
            os.path.realpath(paths["primary"]) == os.path.realpath(paths["backup"]):
        raise HTTPException(400, "Primær og backup er den samme læser – vælg to forskellige")
    db.set_setting("readers", paths)
    if b.get("dedup_sec") is not None:
        db.set_setting("dedup_sec", float(b["dedup_sec"]))
    reader_mgr.configure(paths)
    return {"ok": True}


# ------------------------------------------------------------------ Pi'ens system (kode, førstegangskode)

SYSTEM_HELPER = "/usr/local/sbin/racetimer-system"


def _system(*args, stdin=None):
    """Kald system-hjælperen (root via sudo). RACETIMER_SYSTEM_CMD kan overstyre ved test."""
    import shlex
    import subprocess
    cmd = shlex.split(os.environ.get("RACETIMER_SYSTEM_CMD", "")) or ["sudo", "-n", SYSTEM_HELPER]
    if not os.environ.get("RACETIMER_SYSTEM_CMD") and not os.path.exists(SYSTEM_HELPER):
        return None  # fx under udvikling på en Mac
    r = subprocess.run(cmd + list(args), input=stdin, capture_output=True, text=True, timeout=30)
    try:
        out = json.loads(r.stdout or "{}")
    except ValueError:
        out = {"ok": False, "error": (r.stderr or r.stdout or "Ukendt fejl").strip()[:300]}
    if r.returncode != 0 or out.get("ok") is False:
        raise HTTPException(400, out.get("error") or "Fejl i system-hjælperen")
    return out


def sync_bootcode():
    """Læg førstegangskoden på SD-kortet (til dem uden skærm) – eller fjern den, når koden er skiftet."""
    try:
        _system("bootcode")
    except Exception:  # noqa: BLE001 – må aldrig stoppe RaceTimer
        pass


@app.get("/api/system/pi")
def pi_status(u=Depends(admin)):
    """Har Pi-brugeren stadig standardkoden fra imaget?"""
    try:
        st = _system("status")
    except HTTPException as e:
        return {"available": False, "error": e.detail}
    return {"available": True, **st} if st else {"available": False}


@app.post("/api/system/pi/password")
async def pi_password(request: Request, u=Depends(admin)):
    pw = (await request.json()).get("password") or ""
    if len(pw) < db.MIN_PW:
        raise HTTPException(400, f"Koden skal være mindst {db.MIN_PW} tegn")
    if is_remote(request):
        raise HTTPException(403, "Pi'ens kode kan kun skiftes lokalt")
    out = _system("passwd", stdin=json.dumps({"password": pw}))
    if out is None:
        raise HTTPException(503, "Kan kun skiftes på en Pi installeret fra RaceTimer-imaget")
    return {"ok": True, "user": out.get("user")}


# ------------------------------------------------------------------ WiFi (op til 5 netværk, Pi'en)

WIFI_HELPER = "/usr/local/sbin/racetimer-wifi"


def _wifi(*args, stdin=None):
    """Kald WiFi-hjælperen (root via sudo). RACETIMER_WIFI_CMD kan overstyre ved test."""
    import shlex
    import subprocess
    cmd = shlex.split(os.environ.get("RACETIMER_WIFI_CMD", "")) or ["sudo", "-n", WIFI_HELPER]
    if not os.environ.get("RACETIMER_WIFI_CMD") and not os.path.exists(WIFI_HELPER):
        raise HTTPException(503, "WiFi-styring findes kun på Pi'en (hjælperen er ikke installeret – kør opdateringen)")
    try:
        r = subprocess.run(cmd + list(args), input=stdin, capture_output=True, text=True, timeout=60)
    except subprocess.TimeoutExpired:
        raise HTTPException(504, "WiFi svarede ikke i tide") from None
    try:
        out = json.loads(r.stdout or "{}")
    except ValueError:
        raise HTTPException(500, (r.stderr or r.stdout or "Ukendt fejl fra WiFi-hjælperen").strip()[:300]) from None
    if r.returncode != 0 or out.get("ok") is False:
        raise HTTPException(400, out.get("error") or (r.stderr or "WiFi-fejl").strip()[:300])
    return out


@app.get("/api/wifi")
def wifi_status(u=Depends(admin)):
    return _wifi("status")


@app.put("/api/wifi")
async def wifi_save(request: Request, u=Depends(admin)):
    b = await request.json()
    return _wifi("apply", stdin=json.dumps({"networks": b.get("networks") or []}))


@app.post("/api/wifi/connect/{slot}")
def wifi_connect(slot: int, u=Depends(admin)):
    if not 1 <= slot <= 5:
        raise HTTPException(400, "Ugyldig plads")
    return _wifi("connect", str(slot))


# ------------------------------------------------------------------ fjernadgang (Cloudflare quick tunnel)

TUNNEL_HELPER = "/usr/local/sbin/racetimer-tunnel"
TUNNEL_METRICS = os.environ.get("RACETIMER_TUNNEL_METRICS", "http://127.0.0.1:20241")
TUNNEL_LOG = "/run/racetimer-tunnel/cloudflared.log"


def tunnel_url():
    """Den aktuelle offentlige adresse (https://….trycloudflare.com) – eller None."""
    import re
    import urllib.request
    try:
        with urllib.request.urlopen(TUNNEL_METRICS + "/quicktunnel", timeout=2) as r:
            host = json.loads(r.read() or b"{}").get("hostname")
            if host:
                return "https://" + host
    except (OSError, ValueError):
        pass
    try:  # reserve: læs adressen i cloudflareds log
        with open(TUNNEL_LOG, errors="ignore") as f:
            found = re.findall(r"https://[a-z0-9-]+\.trycloudflare\.com", f.read())
            return found[-1] if found else None
    except OSError:
        return None


def _tunnel(cmd):
    import shlex
    import subprocess
    base = shlex.split(os.environ.get("RACETIMER_TUNNEL_CMD", "")) or ["sudo", "-n", TUNNEL_HELPER]
    if not os.environ.get("RACETIMER_TUNNEL_CMD") and not os.path.exists(TUNNEL_HELPER):
        return {"ok": True, "installed": False, "enabled": False, "active": False}
    try:
        r = subprocess.run(base + [cmd], capture_output=True, text=True, timeout=60)
        return json.loads(r.stdout or "{}")
    except (OSError, ValueError, subprocess.SubprocessError) as e:
        raise HTTPException(500, f"Fjernadgang: {e}") from None


@app.get("/api/remote")
def remote_status(request: Request, u=Depends(admin)):
    st = _tunnel("status")
    st["url"] = tunnel_url() if st.get("active") else None
    st["via_remote"] = is_remote(request)
    st["default_password"] = db.default_password_in_use()
    return st


@app.post("/api/remote")
async def remote_toggle(request: Request, u=Depends(admin)):
    on = bool((await request.json()).get("on"))
    if on and db.default_password_in_use():
        raise HTTPException(400, "Skift førstegangskoden under ⚙ Opsætning → Brugere, før det offentlige link slås til")
    weak = [w["username"] for w in db.q("SELECT username FROM users WHERE strong=0 ORDER BY username")]
    if on and weak:
        raise HTTPException(400, f"Alle administratorer skal have en kode på mindst {db.MIN_PW} tegn, før det offentlige "
                                 f"link slås til. Skift koden for: {', '.join(weak)} (⚙ Opsætning → Brugere)")
    if not on and is_remote(request):
        raise HTTPException(400, "Du er selv logget ind via fjernadgang – slå den fra lokalt på stævnet")
    st = _tunnel("on" if on else "off")
    _remote_cache["t"] = 0
    sync_mgr.poke_tunnel()
    return st


_remote_cache = {"t": 0, "v": None}


def announced_url():
    """Adressen der meldes til hjemmesiden: kun når fjernadgang er slået til (tjekkes højst hvert 30. sek.)."""
    if time.time() - _remote_cache["t"] > 30:
        try:
            on = _tunnel("status").get("active")
        except HTTPException:
            on = False
        _remote_cache.update(t=time.time(), v=tunnel_url() if on else None)
    return _remote_cache["v"]


sync_mgr.tunnel_url = announced_url


# ------------------------------------------------------------------ privat fjernadgang (Tailscale + Raspberry Pi Connect)

REMOTE_HELPER = "/usr/local/sbin/racetimer-remote"


def _remote(*args):
    import shlex
    import subprocess
    base = shlex.split(os.environ.get("RACETIMER_REMOTE_CMD", "")) or ["sudo", "-n", REMOTE_HELPER]
    if not os.environ.get("RACETIMER_REMOTE_CMD") and not os.path.exists(REMOTE_HELPER):
        return {"ok": True, "tailscale": {"installed": False}, "connect": {"installed": False}}
    try:
        r = subprocess.run(base + list(args), capture_output=True, text=True, timeout=45)
        out = json.loads(r.stdout or "{}")
    except (OSError, ValueError, subprocess.SubprocessError) as e:
        raise HTTPException(500, f"Fjernadgang: {e}") from None
    if out.get("ok") is False:
        raise HTTPException(400, out.get("error") or "Fejl")
    return out


@app.get("/api/remote/private")
def private_status(u=Depends(admin)):
    return _remote("status")


@app.post("/api/remote/private")
async def private_action(request: Request, u=Depends(admin)):
    b = await request.json()
    service, action = b.get("service"), b.get("action")
    allowed = {"tailscale": ("up", "down", "logout"), "connect": ("signin", "on", "off", "signout")}
    if action not in allowed.get(service, ()):
        raise HTTPException(400, "Ukendt handling")
    return _remote(service, action)


# ------------------------------------------------------------------ udseende (navn, logo, splash)

# ------------------------------------------------------------------ moduler (eksterne enheder)

def _mod(fn, *a):
    try:
        return fn(*a)
    except modules.ModuleError as e:
        raise HTTPException(400, str(e)) from None


@app.get("/api/modules")
def modules_overview(u=Depends(admin)):
    return module_mgr.overview()


@app.post("/api/modules")
async def modules_add(request: Request, u=Depends(admin)):
    b = await request.json()
    return _mod(module_mgr.add, b.get("kind"), str(b.get("device") or ""), b.get("name"))


@app.put("/api/modules/{mid}")
async def modules_update(mid: str, request: Request, u=Depends(admin)):
    return _mod(module_mgr.update, mid, await request.json())


@app.delete("/api/modules/{mid}")
def modules_remove(mid: str, u=Depends(admin)):
    module_mgr.remove(mid)
    return {"ok": True}


@app.post("/api/modules/{mid}/test")
def modules_test(mid: str, u=Depends(admin)):
    try:
        return {"ok": True, "text": _mod(module_mgr.test, mid)}
    except RuntimeError as e:
        raise HTTPException(502, str(e)) from None


@app.get("/api/branding")
def get_branding():
    pub = branding.public()
    pub["splash_options"] = [{"id": "builtin:" + k, "label": v["label"],
                              "url": f"/static/splash/{k}.webp" if v["file"] else None}
                             for k, v in branding.BUILTIN_SPLASH.items()]
    if branding.get()["splash"].startswith("upload:"):
        pub["splash_options"].append({"id": branding.get()["splash"], "label": "Eget billede", "url": pub["splash"]})
    return pub


@app.get("/branding/logo")
def branding_logo():
    p = branding.logo_path() or _404()
    return FileResponse(p)


@app.get("/branding/splash")
def branding_splash():
    return FileResponse(branding.splash_path())


@app.put("/api/branding")
async def put_branding(request: Request, u=Depends(admin)):
    b = await request.json()
    changes = {}
    if "name" in b:
        name = (b.get("name") or "").strip()
        if not 1 <= len(name) <= 60:
            raise HTTPException(400, "Navnet skal være 1–60 tegn")
        changes["name"] = name
    if "splash" in b:
        choice = b.get("splash") or ""
        kind, _, ref = choice.partition(":")
        ok = (kind == "builtin" and ref in branding.BUILTIN_SPLASH) or \
             (kind == "upload" and ref.startswith("splash.") and branding.safe_asset(ref))
        if not ok:
            raise HTTPException(400, "Ukendt splash")
        changes["splash"] = choice
    if "splash_text" in b:
        text = "\n".join(line.strip() for line in str(b.get("splash_text") or "").strip().splitlines())
        if len(text) > branding.SPLASH_TEXT_MAX:
            raise HTTPException(400, f"Splash-teksten må højst være {branding.SPLASH_TEXT_MAX} tegn")
        changes["splash_text"] = text
    if "scan_banner" in b:
        changes["scan_banner"] = bool(b["scan_banner"])
    out = branding.save(**changes)
    sync_mgr.poke_state()
    return out


@app.post("/api/branding/{kind}")
async def upload_branding(kind: str, file: UploadFile = File(...), u=Depends(admin)):
    if kind not in ("logo", "splash"):
        _404()
    try:
        name = branding.store_upload(kind, file.filename, await file.read())
    except ValueError as e:
        raise HTTPException(400, str(e)) from None
    out = branding.save(**({"logo": name} if kind == "logo" else {"splash": "upload:" + name}))
    sync_mgr.poke_state()
    return out


@app.delete("/api/branding/logo")
def delete_logo(u=Depends(admin)):
    out = branding.save(logo=None)
    sync_mgr.poke_state()
    return out


# ------------------------------------------------------------------ opdatering

@app.get("/api/update")
def update_status(force: bool = False, u=Depends(admin)):
    info = updater.check(force=force)
    try:
        st = updater.helper("status")
        info.update(status=st.get("status"), backups=st.get("backups"), installed=True)
    except RuntimeError as e:
        info.update(installed=False, helper_error=str(e))
    race = db.active_race()
    info["race_running"] = bool(race and race["status"] == "running")
    return info


@app.post("/api/update/install")
async def update_install(request: Request, u=Depends(admin)):
    b = await request.json() if await request.body() else {}
    race = db.active_race()
    if race and race["status"] == "running" and not b.get("force"):
        raise HTTPException(409, f"Løbet “{race['name']}” kører – opdatér efter løbet (eller bekræft for at fortsætte)")
    info = updater.check(force=True)
    if not info.get("available"):
        raise HTTPException(400, "Der er ingen nyere version")
    try:
        tar, sig = updater.download(info)
        return updater.helper("start", tar, sig)
    except (RuntimeError, OSError) as e:
        raise HTTPException(400, str(e)) from None


@app.post("/api/update/rollback")
def update_rollback(u=Depends(admin)):
    try:
        return updater.helper("rollback")
    except RuntimeError as e:
        raise HTTPException(400, str(e)) from None


# ------------------------------------------------------------------ brugere

@app.get("/api/users")
def users(u=Depends(admin)):
    return db.q("SELECT id, username, created_at, strong FROM users ORDER BY username")


@app.post("/api/users")
async def add_user(request: Request, u=Depends(admin)):
    b = await request.json()
    if not b.get("username") or len(b.get("password") or "") < db.MIN_PW:
        raise HTTPException(400, f"Brugernavn og kode på mindst {db.MIN_PW} tegn kræves")
    if db.q1("SELECT 1 FROM users WHERE username=?", (b["username"],)):
        raise HTTPException(409, "Brugernavnet findes allerede")
    return {"id": db.create_user(b["username"], b["password"])}


@app.put("/api/users/{uid}/password")
async def change_pw(uid: int, request: Request, u=Depends(admin)):
    pw = (await request.json()).get("password") or ""
    if len(pw) < db.MIN_PW:
        raise HTTPException(400, f"Koden skal være mindst {db.MIN_PW} tegn")
    if pw == "admin" or pw == db.first_code():
        raise HTTPException(400, "Vælg en ny kode")
    db.set_password(uid, pw)
    threading.Thread(target=sync_bootcode, daemon=True).start()  # fjern førstegangskoden fra SD-kortet
    return {"ok": True}


@app.delete("/api/users/{uid}")
def delete_user(uid: int, u=Depends(admin)):
    if uid == u["id"]:
        raise HTTPException(400, "Du kan ikke slette dig selv")
    if db.q1("SELECT COUNT(*) n FROM users")["n"] <= 1:
        raise HTTPException(400, "Der skal være mindst én admin")
    db.ex("DELETE FROM users WHERE id=?", (uid,))
    return {"ok": True}
